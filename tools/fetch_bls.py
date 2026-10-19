#!/usr/bin/env python3
"""Lädt den Bundeslebensmittelschlüssel (BLS 4.0) und erzeugt data/bls.json für die App.

Quelle: Max Rubner-Institut (2025): Bundeslebensmittelschlüssel (BLS), Version 4.0.
Karlsruhe. Lizenz CC BY 4.0 – https://blsdb.de

Aufruf:
  python tools/fetch_bls.py data/bls.json            # lädt von blsdb.de
  python tools/fetch_bls.py data/bls.json datei.zip  # nutzt eine bereits geladene ZIP/XLSX

Schlägt der Download fehl, endet das Skript ohne Fehler (die App funktioniert auch ohne).
"""
import io
import json
import re
import sys
import urllib.request
import zipfile

DOWNLOAD_PAGE = "https://blsdb.de/download"
UA = {"User-Agent": "Mozilla/5.0 (Formkurve BLS import)"}
CREDIT = "Max Rubner-Institut (2025): Bundeslebensmittelschlüssel (BLS), Version 4.0. Karlsruhe. Lizenz CC BY 4.0"


def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120) as r:
        return r.read()


def find_zip_url():
    html = get(DOWNLOAD_PAGE).decode("utf-8", "replace")
    links = re.findall(r'href="([^"]+\.zip[^"]*)"', html)
    de = [l for l in links if "_DE" in l] or links
    if not de:
        raise RuntimeError("Kein ZIP-Link auf der Download-Seite gefunden")
    url = de[0].replace("&amp;", "&")
    if url.startswith("/"):
        url = "https://blsdb.de" + url
    return url


def xlsx_from(data, name):
    if name.lower().endswith(".zip") or data[:2] == b"PK" and not name.lower().endswith(".xlsx"):
        z = zipfile.ZipFile(io.BytesIO(data))
        names = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        pref = [n for n in names if "daten" in n.lower()] or names
        if not pref:
            raise RuntimeError("Keine XLSX im ZIP")
        return z.read(pref[0])
    return data


def detect(header):
    h = [str(x or "").strip() for x in header]

    def find(*pats):
        for p in pats:
            for i, x in enumerate(h):
                if re.search(p, x, re.I):
                    return i
        return -1

    return {
        "code": find(r"^(SBLS|BLS[ _-]?Code|Code)\b", r"schl(ue|ü)ssel"),
        "name": find(r"bezeichnung(?!.*engl)", r"^name(?!.*engl)", r"lebensmittel(?!.*engl)"),
        "kcal": find(r"^ENERCC\b", r"kcal"),
        "protein": find(r"^PROT", r"protein|eiwei"),
        "fat": find(r"^FAT\b", r"^fett\b"),
        "carbs": find(r"^CHO\b", r"kohlenhydrate.*verf", r"kohlenhydrat"),
    }


def num(v):
    if v is None or v == "":
        return 0
    try:
        return round(float(str(v).replace(" ", "").replace(",", ".")), 1)
    except ValueError:
        return 0


def convert(xlsx_bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    best = max((list(ws.iter_rows(values_only=True)) for ws in wb.worksheets), key=len)
    hi = next((i for i, r in enumerate(best[:10]) if detect(r)["kcal"] >= 0 and detect(r)["name"] >= 0), -1)
    if hi < 0:
        raise RuntimeError("Spalten für Name/Kalorien nicht gefunden: %r" % (best[0][:12],))
    c = detect(best[hi])
    items = []
    for r in best[hi + 1:]:
        name = str(r[c["name"]] or "").strip()
        if not name:
            continue
        code = str(r[c["code"]] or "") if c["code"] >= 0 else str(len(items))
        items.append([code, name, num(r[c["kcal"]]), num(r[c["protein"]] if c["protein"] >= 0 else 0),
                      num(r[c["carbs"]] if c["carbs"] >= 0 else 0), num(r[c["fat"]] if c["fat"] >= 0 else 0)])
    return items


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "data/bls.json"
    try:
        if len(sys.argv) > 2:
            src = sys.argv[2]
            data = open(src, "rb").read()
        else:
            src = find_zip_url()
            print("Lade", src.split("?")[0])
            data = get(src)
        items = convert(xlsx_from(data, src.split("?")[0]))
    except Exception as e:  # noqa: BLE001 – Import ist optional
        print("BLS-Import übersprungen:", e)
        return 0
    with open(out, "w", encoding="utf-8") as f:
        json.dump({"source": CREDIT, "items": items}, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{len(items)} Lebensmittel nach {out} geschrieben")
    return 0


if __name__ == "__main__":
    sys.exit(main())
